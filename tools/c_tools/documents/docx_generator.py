"""
Word (.docx) Document Generator module.
Generates structured Word documents with headings, paragraphs, bullet lists, and tables.
Supports python-docx with clean HTML/Markdown fallback.
"""

import time
from pathlib import Path
from typing import Dict, Any, Optional, List

from tools.c_tools.schemas import ExecutionResult, ExecutionError, ErrorCode
from tools.c_tools.file_ops.path_security import validate_workspace_path


def generate_docx(
    filename: str,
    title: str,
    sections: List[Dict[str, Any]],
    workspace_root: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Generates a Word (.docx) document with structured headings, text, bullet lists, and tables.

    Args:
        filename: Target output filename (e.g., 'report.docx').
        title: Main document title.
        sections: List of section dicts. Each dict can contain 'heading', 'content', 'bullet_points', 'table'.
        workspace_root: Optional workspace root directory override.

    Returns:
        Structured ExecutionResult dictionary.
    """
    start_time = time.perf_counter()

    if not filename.endswith(".docx"):
        filename += ".docx"

    try:
        validated_path = validate_workspace_path(filename, workspace_root)
    except ValueError as e:
        return ExecutionResult(
            success=False,
            execution_time=time.perf_counter() - start_time,
            error=ExecutionError(code=ErrorCode.PATH_TRAVERSAL_ATTEMPT, message=str(e)),
        ).to_dict()

    validated_path.parent.mkdir(parents=True, exist_ok=True)

    # Check python-docx availability
    try:
        import docx
        from docx import Document

        doc = Document()
        doc.add_heading(title, level=0)

        for sec in sections:
            if "heading" in sec:
                doc.add_heading(str(sec["heading"]), level=1)

            if "content" in sec:
                content_val = sec["content"]
                if isinstance(content_val, list):
                    for p in content_val:
                        doc.add_paragraph(str(p))
                else:
                    doc.add_paragraph(str(content_val))

            if "bullet_points" in sec and isinstance(sec["bullet_points"], list):
                for bullet in sec["bullet_points"]:
                    doc.add_paragraph(str(bullet), style="List Bullet")

            if "table" in sec and isinstance(sec["table"], dict):
                tbl_data = sec["table"]
                headers = tbl_data.get("headers", [])
                rows = tbl_data.get("rows", [])
                
                if headers or rows:
                    cols_cnt = len(headers) if headers else (len(rows[0]) if rows else 1)
                    table = doc.add_table(rows=0, cols=cols_cnt)
                    
                    if headers:
                        hdr_cells = table.add_row().cells
                        for i, header_text in enumerate(headers):
                            hdr_cells[i].text = str(header_text)
                    
                    for r in rows:
                        row_cells = table.add_row().cells
                        for i, val in enumerate(r):
                            if i < cols_cnt:
                                row_cells[i].text = str(val)

        doc.save(str(validated_path))
        format_used = "native_docx"

    except ImportError:
        # Fallback to structured text/HTML formatting if python-docx is absent
        fallback_lines = [f"# {title}\n"]
        for sec in sections:
            if "heading" in sec:
                fallback_lines.append(f"\n## {sec['heading']}\n")
            if "content" in sec:
                fallback_lines.append(f"{sec['content']}\n")
            if "bullet_points" in sec:
                for b in sec["bullet_points"]:
                    fallback_lines.append(f"* {b}")
                fallback_lines.append("")
            if "table" in sec:
                tbl = sec["table"]
                headers = tbl.get("headers", [])
                if headers:
                    fallback_lines.append("| " + " | ".join(headers) + " |")
                    fallback_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                for r in tbl.get("rows", []):
                    fallback_lines.append("| " + " | ".join(str(cell) for cell in r) + " |")
                fallback_lines.append("")

        fallback_lines.append("\n<!-- Note: Install python-docx via 'pip install python-docx' for native .docx binary generation -->\n")
        
        with open(validated_path, "w", encoding="utf-8") as f:
            f.write("\n".join(fallback_lines))
        format_used = "markdown_fallback"

    except Exception as e:
        return ExecutionResult(
            success=False,
            execution_time=time.perf_counter() - start_time,
            error=ExecutionError(
                code=ErrorCode.DOCUMENT_GENERATION_ERROR,
                message=f"Failed to generate Word document '{filename}': {str(e)}",
            ),
        ).to_dict()

    exec_time = time.perf_counter() - start_time
    file_size = validated_path.stat().st_size

    return ExecutionResult(
        success=True,
        result={
            "filename": validated_path.name,
            "path": str(validated_path),
            "size_bytes": file_size,
            "sections_count": len(sections),
            "format_used": format_used,
        },
        execution_time=exec_time,
    ).to_dict()
